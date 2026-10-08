"""
Step 1: 扫描PPT中所有图片占位符
输出 placeholders.json
"""
import sys
import re
import json
from pathlib import Path
from pptx import Presentation

# 占位符识别正则
PLACEHOLDER_PATTERNS = [
    r'\[\s*图片占位\s*\]',
    r'\[\s*占位图\s*\d*\s*\]',
    r'请提供.{0,30}图',
    r'\[\s*placeholder\s*\]',
    r'\[\s*PLACEHOLDER\s*\]',
    r'TODO.{0,30}图',
    r'\[\s*image\s+placeholder\s*\]',
]


def is_placeholder(text: str) -> bool:
    """判断文本框是否为占位符"""
    if not text or len(text) > 500:
        return False
    return any(re.search(p, text, re.IGNORECASE) for p in PLACEHOLDER_PATTERNS)


def clean_description(text: str) -> str:
    """清理占位符文本，提取真实描述"""
    text = text.strip()
    for p in PLACEHOLDER_PATTERNS:
        text = re.sub(p, '', text, flags=re.IGNORECASE)
    text = re.sub(r'\s+', ' ', text).strip()
    return text


def scan(pptx_path: str) -> list:
    """扫描PPT，返回所有占位符列表"""
    prs = Presentation(pptx_path)
    placeholders = []

    for slide_idx, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text
            if not is_placeholder(text):
                continue

            description = clean_description(text)
            if not description:
                description = f"slide {slide_idx} placeholder"

            placeholders.append({
                "slide_idx": slide_idx,
                "shape_id": shape.shape_id,
                "shape_name": shape.name,
                "bbox": {
                    "left": int(shape.left),
                    "top": int(shape.top),
                    "width": int(shape.width),
                    "height": int(shape.height),
                },
                "raw_text": text,
                "description": description,
                "category": None,
            })

    return placeholders


def main():
    if len(sys.argv) < 2:
        print("用法: python 01_scan_placeholders.py <input.pptx>")
        sys.exit(1)

    pptx_path = sys.argv[1]
    if not Path(pptx_path).exists():
        print(f"错误: 文件不存在 {pptx_path}")
        sys.exit(1)

    print(f"扫描 {pptx_path} ...")
    placeholders = scan(pptx_path)

    out_path = Path("placeholders.json")
    out_path.write_text(
        json.dumps(placeholders, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )

    print(f"扫描完成: 共 {len(placeholders)} 个占位符")
    print(f"输出: {out_path.absolute()}")

    # 按slide分组打印
    by_slide = {}
    for ph in placeholders:
        by_slide.setdefault(ph["slide_idx"], []).append(ph)

    for slide_idx in sorted(by_slide.keys()):
        items = by_slide[slide_idx]
        print(f"  Slide {slide_idx}: {len(items)} 个占位符")
        for ph in items:
            desc_preview = ph["description"][:50].replace("\n", " ")
            print(f"    - {desc_preview}...")


if __name__ == "__main__":
    main()
