"""
Step 5: 将generated/目录里的图片按bbox坐标回写到PPT
保留原PPT其他所有内容不动
"""
import sys
import json
from pathlib import Path
from pptx import Presentation


def inject(pptx_path: str, placeholders_json: str, output_path: str):
    """回写生成的图片到PPT"""
    prs = Presentation(pptx_path)
    placeholders = json.loads(Path(placeholders_json).read_text(encoding="utf-8"))
    generated_dir = Path("generated")

    if not generated_dir.exists():
        print("错误: 未找到 generated/ 目录")
        return

    # 按slide_idx分组
    by_slide = {}
    for ph in placeholders:
        by_slide.setdefault(ph["slide_idx"], []).append(ph)

    total_injected = 0
    total_skipped = 0

    for slide_idx, slide in enumerate(prs.slides, 1):
        if slide_idx not in by_slide:
            continue

        slide_placeholders = by_slide[slide_idx]
        print(f"Slide {slide_idx}: 处理 {len(slide_placeholders)} 个占位符")

        for ph in slide_placeholders:
            img_filename = f"slide{slide_idx}_shape{ph['shape_id']}.png"
            img_path = generated_dir / img_filename

            if not img_path.exists():
                print(f"  跳过 shape {ph['shape_id']}: 未找到 {img_filename}")
                total_skipped += 1
                continue

            bbox = ph["bbox"]

            # 找到原占位符shape并删除
            target_shape = None
            for shape in slide.shapes:
                if shape.shape_id == ph["shape_id"]:
                    target_shape = shape
                    break

            if target_shape is not None:
                sp = target_shape._element
                sp.getparent().remove(sp)
                print(f"  删除原占位符 shape_id={ph['shape_id']}")
            else:
                print(f"  警告: 未找到原shape_id={ph['shape_id']}, 直接添加图片")

            # 在原坐标插入图片
            slide.shapes.add_picture(
                str(img_path),
                left=bbox["left"],
                top=bbox["top"],
                width=bbox["width"],
                height=bbox["height"],
            )
            print(f"  ✓ 插入图片: {img_filename}")
            total_injected += 1

    prs.save(output_path)
    print(f"\n{'=' * 50}")
    print(f"完成: 注入 {total_injected} 张图片, 跳过 {total_skipped} 个")
    print(f"输出: {output_path}")


def main():
    if len(sys.argv) < 2:
        print("用法: python 06_inject_to_pptx.py <input.pptx>")
        sys.exit(1)

    input_path = sys.argv[1]
    if not Path(input_path).exists():
        print(f"错误: 文件不存在 {input_path}")
        sys.exit(1)

    placeholders_path = Path("placeholders.json")
    if not placeholders_path.exists():
        print("错误: 未找到 placeholders.json")
        sys.exit(1)

    output_path = input_path.replace(".pptx", "_polished.pptx")
    inject(input_path, str(placeholders_path), output_path)


if __name__ == "__main__":
    main()
